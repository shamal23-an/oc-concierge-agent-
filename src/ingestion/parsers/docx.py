from __future__ import annotations

from pathlib import Path

import structlog

logger = structlog.get_logger()


class DocxParser:
    """DOCX parser using python-docx."""

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() in {".docx", ".doc"}

    def parse(self, path: Path) -> str | None:
        try:
            from docx import Document

            doc = Document(str(path))
            parts: list[str] = []

            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text.strip())

            for table in doc.tables:
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    parts.append("\n".join(rows))

            text = "\n\n".join(parts)
            if len(text.strip()) < 20:
                logger.warning("docx_parse_empty", path=str(path))
                return None
            return text.strip()
        except Exception:
            logger.warning("docx_parse_failed", path=str(path))
            return None
